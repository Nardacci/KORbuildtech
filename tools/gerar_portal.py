"""Gera o portal de especificações e manual do KORbuild (portal/) a partir do Markdown em portal/fonte/.

Uso:
  pip install markdown python-docx
  node tools/capturas-portal.mjs          # (opcional) refaz as capturas das telas em portal/img/
  python3 tools/gerar_portal.py           # gera as páginas (inglês em portal/, português em portal/pt/), a busca, os Word e confere os links

O texto é escrito uma vez só (portal/fonte/<idioma>/...). Daqui saem a página web e o Word de cada documento.
Markdown aceito, além do comum (títulos, listas, tabelas, negrito, links):
  ![Legenda](img/arquivo.png)        imagem com legenda (arquivo terminado em -mobile.png sai no tamanho de celular)
  [texto](@/daily/user-manual.html)   link para outra página do portal ("@/" é a raiz do portal)
  !!! who "Who can do it"            quadros: who (quem pode), tip, warning, simulated (simulado no protótipo),
                                     confirm (a confirmar), say (o que dizer na apresentação), note
"""
import html
import json
import os
import re
import sys
from datetime import date
from html.parser import HTMLParser
from pathlib import Path

import markdown

RAIZ = Path(__file__).resolve().parent.parent
PORTAL = RAIZ / 'portal'
FONTE = PORTAL / 'fonte'
# Dois idiomas: inglês na raiz (portal/) e português em portal/pt/, cada um com as suas capturas (img/).
IDIOMAS = ['en', 'pt']
IDIOMA = 'en'
SAIDA = PORTAL      # onde as páginas do idioma são gravadas
TOPO = ''           # caminho de SAIDA até portal/ ('' no inglês, '../' no português)
MESES = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']


def hoje():
    d = date.today()
    return f'{d.day} de {MESES[d.month - 1]} de {d.year}' if IDIOMA == 'pt' else d.strftime('%B %-d, %Y')


def definir_idioma(idioma):
    global IDIOMA, SAIDA, TOPO
    IDIOMA = idioma
    SAIDA = PORTAL if idioma == 'en' else PORTAL / idioma
    TOPO = '' if idioma == 'en' else '../'


# Textos da moldura (o conteúdo vem do Markdown de cada idioma). Chave = o texto em inglês.
PT = {
    'Daily construction reports from the field: crew, weather, work performed, GPS-stamped photos and issues. The office reviews, approves and shares a sealed PDF with the client. Also home to the job list and the job schedule.':
        'O diário de obra feito no canteiro: equipe, clima, serviços executados, fotos com GPS e ocorrências. A aprovação lacra o relatório e ele é compartilhado com o cliente em PDF. Aqui também ficam as obras e o cronograma.',
    'Crew time clock with GPS and geofence, weekly timesheets with overtime, approval and payroll export, and labor cost by job: budget × actual × forecast to completion.':
        'Ponto da equipe com GPS e cerca, timesheets semanais com hora extra, aprovação e exportação para a folha, e custo de mão de obra por obra: orçado × realizado × projeção.',
    'Plan takeoff in feet and inches: scale, measure, assemblies, quantities, quote lists and proposals.':
        'Medição de plantas em pés e polegadas: escala, medição, assemblies, quantidades, lista para cotação e proposta.',
    'Company setup: employees, labor burden, work rules, jobs and their schedules, contacts, users and access roles, audit log.':
        'Configuração da empresa: funcionários, encargos, regras de jornada, obras e cronogramas, contatos, usuários e perfis de acesso, auditoria.',
    'Bids': 'Concorrências',
    'Invitations to bid between companies: the contracting company sends the bid package, levels the bids and awards; the subcontractor bids from Measure or by a link; a channel for questions that becomes the job conversation; subcontracting down the chain.':
        'Concorrências entre empresas: a contratante manda o pacote, compara as propostas e fecha; a prestadora cota pelo Measure ou por um link; um canal de perguntas que vira a conversa da obra; e a quarteirização.',
    'How a bid works, who takes part and how it links the companies on the job.': 'Como funciona uma concorrência, quem participa e como ela liga as empresas na obra.',
    'Send a bid, level and award; bid as a subcontractor; the channel; subcontracting — step by step.': 'Mandar a concorrência, comparar e fechar; cotar como prestadora; o canal; a quarteirização — passo a passo.',
    'Who sees what, revisions, due dates, leveling, award and subcontracting rules, with IDs.': 'Quem vê o quê, revisões, prazos, comparação, fechamento e quarteirização, com código.',
    'Who runs bids and who answers invitations.': 'Quem faz concorrências e quem responde convites.',
    'Overview': 'Visão geral', 'User manual': 'Manual do usuário', 'Business rules': 'Regras de negócio',
    'Roles and permissions': 'Perfis e permissões', 'Prototype status': 'Situação do protótipo',
    'Start here': 'Comece aqui', 'Using it': 'No dia a dia', 'Rules and access': 'Regras e acesso',
    'What Daily does, who uses it and the main screens.': 'O que o Daily faz, quem usa e as telas principais.',
    'Step by step, task by task, with screenshots.': 'Passo a passo, tarefa por tarefa, com as telas.',
    'Every rule the system enforces, with an ID.': 'Cada regra que o sistema aplica, com um código.',
    'Who can do what in Daily.': 'Quem pode fazer o quê no Daily.',
    'What is simulated today and what comes next.': 'O que é simulado hoje e o que vem depois.',
    'What Crew does, who uses it and the main screens.': 'O que o Crew faz, quem usa e as telas principais.',
    'Clock in, approve timesheets, follow labor cost — step by step.': 'Bater o ponto, aprovar timesheets, acompanhar o custo — passo a passo.',
    'Punches, geofence, overtime, cost and forecast rules, with IDs.': 'Regras de batidas, cerca, hora extra, custo e projeção, com códigos.',
    'Who can do what in Crew.': 'Quem pode fazer o quê no Crew.',
    'What Measure does, who uses it and how a takeoff becomes a proposal.': 'O que o Measure faz, quem usa e como a medição vira proposta.',
    'From the PDF plan to the proposal, step by step, with screenshots.': 'Do PDF da planta à proposta, passo a passo, com as telas.',
    'Units, scale, measuring, formulas, pricing, reports and typicals, with IDs.': 'Unidades, escala, medição, fórmulas, preços, relatórios e typicals, com códigos.',
    'Who can measure and who can change the catalog.': 'Quem pode medir e quem pode mudar o catálogo.',
    'What is in Settings and who uses it.': 'O que há no Settings e quem usa.',
    'Employees, rules, company, contacts, jobs, the job schedule, users and roles — step by step.': 'Funcionários, regras, empresa, contatos, obras, cronograma, usuários e perfis — passo a passo.',
    'Effective dates, employees, jobs, schedule and access rules, with IDs.': 'Vigência, funcionários, obras, cronograma e acesso, com códigos.',
    'The full permission catalog and the default roles.': 'O catálogo completo de permissões e os perfis prontos.',
    'What is simulated, what is still fixed and what comes next.': 'O que é simulado, o que ainda é fixo e o que vem depois.',
    'Who the customers are (general contractors and subcontractors), who registers what and what each company sees of the others.':
        'Quem são os clientes (construtoras e prestadoras), quem cadastra o quê e o que cada empresa vê da outra.',
    'How to present KORbuild to a customer: the story, the order of the screens, which demo user to sign in with and what to say at each step.':
        'Como apresentar o KORbuild a um cliente: a história, a ordem das telas, com qual usuário entrar e o que dizer em cada passo.',
    'KORbuild product guide': 'Guia do produto KORbuild',
    'KORbuild product guide: specifications, business rules and user manual for each module.': 'Guia do produto KORbuild: especificações, regras de negócio e manual do usuário de cada módulo.',
    'KORbuild product guide, home': 'Guia do produto KORbuild, início', 'Guide': 'Guia', 'Main': 'Principal',
    'Modules': 'Módulos', 'How it is organized': 'Como se organiza', "Presenter's guide": 'Guia do apresentador',
    'Search rules, screens, tasks…': 'Buscar regras, telas, tarefas…', 'Search the guide': 'Buscar no guia',
    'Open the prototype': 'Abrir o protótipo', 'Menu and search': 'Menu e busca',
    'updated': 'atualizado em', 'Main sign-in': 'Entrada principal', 'Prototype': 'Protótipo',
    'Screenshots come from the prototype with sample data. Names, companies and numbers are fictitious.':
        'As telas vêm do protótipo, com dados de exemplo. Nomes, empresas e números são fictícios.',
    'Back to top': 'Voltar ao topo', 'On this page': 'Nesta página', 'You are here': 'Você está aqui', 'Home': 'Início',
    'documents': 'documentos', 'Download Word': 'Baixar Word', 'Print': 'Imprimir', 'KORbuild guide': 'Guia KORbuild',
    'Coming soon': 'Em breve',
    'Documentation coming soon': 'Documentação em breve',
    'Everything KORbuild does, in one place.': 'Tudo o que o KORbuild faz, num lugar só.',
    'KORbuild is a platform for construction companies in the United States: general contractors and the trade contractors they hire. This guide explains each module — what it is for, the rules it enforces, who can do what — and walks through every task with screenshots, so anyone can present the product with confidence.':
        'O KORbuild é uma plataforma para empresas de construção nos Estados Unidos: construtoras e as prestadoras de serviço que elas contratam. Este guia explica cada módulo — para que serve, as regras que aplica, quem pode fazer o quê — e mostra cada tarefa com as telas, para qualquer pessoa apresentar o produto com segurança.',
    'Presenter\'s guide →': 'Guia do apresentador →', 'How KORbuild is organized': 'Como o KORbuild se organiza',
    'Explore the modules': 'Conhecer os módulos', 'Open the prototype ↗': 'Abrir o protótipo ↗',
    'One account, one set of jobs and people, shared by every module. Each company turns on the modules it needs.':
        'Uma conta, um cadastro de obras e de pessoas, usado por todos os módulos. Cada empresa liga os módulos de que precisa.',
    'Where to start': 'Por onde começar', 'Pick the path that matches why you are here.': 'Escolha o caminho que combina com o motivo de estar aqui.',
    'I will present KORbuild to a customer': 'Vou apresentar o KORbuild a um cliente',
    "Read the presenter's guide: the story, the order of the screens and what to say.": 'Leia o guia do apresentador: a história, a ordem das telas e o que dizer.',
    'I will use KORbuild every day': 'Vou usar o KORbuild no dia a dia',
    'Go to your module and open the user manual: one task at a time, with screenshots.': 'Vá ao seu módulo e abra o manual do usuário: uma tarefa por vez, com as telas.',
    'I need the rules and permissions': 'Preciso das regras e das permissões',
    'Each module lists its business rules with IDs and who can do what.': 'Cada módulo lista as regras de negócio, com códigos, e quem pode fazer o quê.',
    'user manual': 'manual do usuário', 'business rules': 'regras de negócio', 'roles and permissions': 'perfis e permissões',
    'The documentation of this module is being written. Meanwhile, see it working in the': 'A documentação deste módulo está sendo escrita. Enquanto isso, veja-o funcionando no',
    'prototype': 'protótipo',
    'The bell, the emails KORbuild sends, how to choose what arrives by email and the rules behind it.': 'O sininho, os e-mails que o KORbuild envia, como escolher o que chega por e-mail e as regras por trás.',
    'The KORbuild team panel: companies, plans, suspension, email queue and default values.': 'O painel da equipe do KORbuild: empresas, planos, suspensão, fila de e-mails e valores padrão.',
}


def L(texto):
    """O texto da moldura no idioma em geração."""
    return texto if IDIOMA == 'en' else PT.get(texto, texto)

# ---------------------------------------------------------------- catálogo

MODULOS = [
    {'id': 'daily', 'nome': 'Daily', 'cor': 'var(--cor-daily)', 'ico': 'D',
     'resumo': 'Daily construction reports from the field: crew, weather, work performed, GPS-stamped photos and issues. '
               'The office reviews, approves and shares a sealed PDF with the client. Also home to the job list and the job schedule.',
     'paginas': [
         {'slug': 'overview', 'titulo': 'Overview', 'grupo': 'Start here', 'desc': 'What Daily does, who uses it and the main screens.'},
         {'slug': 'user-manual', 'titulo': 'User manual', 'grupo': 'Using it', 'desc': 'Step by step, task by task, with screenshots.'},
         {'slug': 'business-rules', 'titulo': 'Business rules', 'grupo': 'Rules and access', 'desc': 'Every rule the system enforces, with an ID.'},
         {'slug': 'roles-and-permissions', 'titulo': 'Roles and permissions', 'grupo': 'Rules and access', 'desc': 'Who can do what in Daily.'},
         {'slug': 'prototype-status', 'titulo': 'Prototype status', 'grupo': 'Rules and access', 'desc': 'What is simulated today and what comes next.'},
     ]},
    {'id': 'crew', 'nome': 'Crew', 'cor': 'var(--cor-crew)', 'ico': 'C',
     'resumo': 'Crew time clock with GPS and geofence, weekly timesheets with overtime, approval and payroll export, '
               'and labor cost by job: budget × actual × forecast to completion.',
     'paginas': [
         {'slug': 'overview', 'titulo': 'Overview', 'grupo': 'Start here', 'desc': 'What Crew does, who uses it and the main screens.'},
         {'slug': 'user-manual', 'titulo': 'User manual', 'grupo': 'Using it', 'desc': 'Clock in, approve timesheets, follow labor cost — step by step.'},
         {'slug': 'business-rules', 'titulo': 'Business rules', 'grupo': 'Rules and access', 'desc': 'Punches, geofence, overtime, cost and forecast rules, with IDs.'},
         {'slug': 'roles-and-permissions', 'titulo': 'Roles and permissions', 'grupo': 'Rules and access', 'desc': 'Who can do what in Crew.'},
         {'slug': 'prototype-status', 'titulo': 'Prototype status', 'grupo': 'Rules and access', 'desc': 'What is simulated today and what comes next.'},
     ]},
    {'id': 'measure', 'nome': 'Measure', 'cor': 'var(--cor-measure)', 'ico': 'M',
     'resumo': 'Plan takeoff in feet and inches: scale, measure, assemblies, quantities, quote lists and proposals.',
     'paginas': [
         {'slug': 'overview', 'titulo': 'Overview', 'grupo': 'Start here', 'desc': 'What Measure does, who uses it and how a takeoff becomes a proposal.'},
         {'slug': 'user-manual', 'titulo': 'User manual', 'grupo': 'Using it', 'desc': 'From the PDF plan to the proposal, step by step, with screenshots.'},
         {'slug': 'business-rules', 'titulo': 'Business rules', 'grupo': 'Rules and access', 'desc': 'Units, scale, measuring, formulas, pricing, reports and typicals, with IDs.'},
         {'slug': 'roles-and-permissions', 'titulo': 'Roles and permissions', 'grupo': 'Rules and access', 'desc': 'Who can measure and who can change the catalog.'},
         {'slug': 'prototype-status', 'titulo': 'Prototype status', 'grupo': 'Rules and access', 'desc': 'What is simulated today and what comes next.'},
     ]},
    {'id': 'settings', 'nome': 'Settings', 'cor': 'var(--cor-settings)', 'ico': 'S',
     'resumo': 'Company setup: employees, labor burden, work rules, jobs and their schedules, contacts, users and access roles, audit log.',
     'paginas': [
         {'slug': 'overview', 'titulo': 'Overview', 'grupo': 'Start here', 'desc': 'What is in Settings and who uses it.'},
         {'slug': 'user-manual', 'titulo': 'User manual', 'grupo': 'Using it', 'desc': 'Employees, rules, company, contacts, jobs, the job schedule, users and roles — step by step.'},
         {'slug': 'business-rules', 'titulo': 'Business rules', 'grupo': 'Rules and access', 'desc': 'Effective dates, employees, jobs, schedule and access rules, with IDs.'},
         {'slug': 'roles-and-permissions', 'titulo': 'Roles and permissions', 'grupo': 'Rules and access', 'desc': 'The full permission catalog and the default roles.'},
         {'slug': 'prototype-status', 'titulo': 'Prototype status', 'grupo': 'Rules and access', 'desc': 'What is simulated, what is still fixed and what comes next.'},
     ]},
    {'id': 'bids', 'nome': 'Bids', 'cor': 'var(--cor-bids)', 'ico': 'B',
     'resumo': 'Invitations to bid between companies: the contracting company sends the bid package, levels the bids and awards; '
               'the subcontractor bids from Measure or by a link; a channel for questions that becomes the job conversation; subcontracting down the chain.',
     'paginas': [
         {'slug': 'overview', 'titulo': 'Overview', 'grupo': 'Start here', 'desc': 'How a bid works, who takes part and how it links the companies on the job.'},
         {'slug': 'user-manual', 'titulo': 'User manual', 'grupo': 'Using it', 'desc': 'Send a bid, level and award; bid as a subcontractor; the channel; subcontracting — step by step.'},
         {'slug': 'business-rules', 'titulo': 'Business rules', 'grupo': 'Rules and access', 'desc': 'Who sees what, revisions, due dates, leveling, award and subcontracting rules, with IDs.'},
         {'slug': 'roles-and-permissions', 'titulo': 'Roles and permissions', 'grupo': 'Rules and access', 'desc': 'Who runs bids and who answers invitations.'},
         {'slug': 'prototype-status', 'titulo': 'Prototype status', 'grupo': 'Rules and access', 'desc': 'What is simulated today and what comes next.'},
     ]},
]
GERAIS = [
    {'slug': 'organization', 'titulo': 'How KORbuild is organized', 'desc': 'Who the customers are (general contractors and subcontractors), '
                                                                      'who registers what and what each company sees of the others.'},
    {'slug': 'notifications', 'titulo': 'Notifications and email', 'desc': 'The bell, the emails KORbuild sends, how to choose what arrives by email and the rules behind it.'},
    {'slug': 'platform-admin', 'titulo': 'Platform administration', 'desc': 'The KORbuild team panel: companies, plans, suspension, email queue and default values.'},
    {'slug': 'presenting', 'titulo': "Presenter's guide", 'desc': 'How to present KORbuild to a customer: the story, the order of the screens, '
                                                                 'which demo user to sign in with and what to say at each step.'},
]
GRUPOS = ['Start here', 'Using it', 'Rules and access']

E = html.escape


# ---------------------------------------------------------------- Markdown → HTML

def converter(md_texto, raiz):
    md = markdown.Markdown(extensions=['tables', 'toc', 'attr_list', 'admonition', 'md_in_html', 'sane_lists'],
                           extension_configs={'toc': {'toc_depth': '2-3'}})
    corpo = md.convert(md_texto)
    toc = md.toc_tokens
    # links "@/..." = raiz do portal
    corpo = corpo.replace('href="@/', 'href="' + raiz)
    # imagem sozinha no parágrafo vira figura com legenda; "-mobile" no nome = tamanho de celular
    def figura(alt, src):
        alt = html.unescape(alt)
        classe = ' class="celular"' if '-mobile' in src else ''
        return f'<figure{classe}><img src="{raiz}{src}" alt="{E(alt)}" loading="lazy"><figcaption>{E(alt)}</figcaption></figure>'
    corpo = re.sub(r'<p><img alt="([^"]*)" src="(img/[^"]+)"\s*/?></p>', lambda m: figura(m.group(1), m.group(2)), corpo)
    return corpo, toc


def titulo_e_corpo(caminho):
    texto = caminho.read_text(encoding='utf-8')
    m = re.match(r'#\s+(.+)\n', texto)
    titulo = m.group(1).strip() if m else caminho.stem
    return titulo, texto[m.end():] if m else texto


# ---------------------------------------------------------------- moldura

def cabecalho(titulo, raiz, ativo='', caminho='index.html'):
    def nav(id_, href, rotulo):
        return f'<a href="{raiz}{href}"{" aria-current=page" if ativo == id_ else ""}>{rotulo}</a>'
    # a mesma página no outro idioma
    outro = 'pt' if IDIOMA == 'en' else 'en'
    href_outro = f'{raiz}pt/{caminho}' if IDIOMA == 'en' else f'{raiz}../{caminho}'
    idiomas = (f'<span class="idioma-topo" aria-label="Language / Idioma">'
               f'<a href="{href_outro if IDIOMA == "pt" else "#"}"{" aria-current=true" if IDIOMA == "en" else ""} lang="en">EN</a>'
               f'<a href="{href_outro if IDIOMA == "en" else "#"}"{" aria-current=true" if IDIOMA == "pt" else ""} lang="pt-BR">PT</a></span>')
    return f"""<!DOCTYPE html>
<html lang="{'pt-BR' if IDIOMA == 'pt' else 'en'}">
<head>
<meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{E(titulo)}</title>
<meta name="description" content="{L('KORbuild product guide: specifications, business rules and user manual for each module.')}">
<link rel="icon" href="{raiz}{TOPO}../icons/icon-192.png">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap" rel="stylesheet">
<link rel="stylesheet" href="{raiz}{TOPO}assets/portal.css">
</head>
<body data-raiz="{raiz}">
<header class="topo">
  <a class="marca" href="{raiz}index.html" aria-label="{L('KORbuild product guide, home')}"><img src="{raiz}{TOPO}../icons/icon-192.png" alt=""><span><b>KOR</b>build</span><small>{L('Guide')}</small></a>
  <nav class="topo-nav" aria-label="{L('Main')}">
    {nav('modulos', 'index.html#modules', L('Modules'))}
    {nav('organization', 'organization.html', L('How it is organized'))}
    {nav('presenting', 'presenting.html', L("Presenter's guide"))}
  </nav>
  <div class="topo-acoes">
    <label class="busca"><span aria-hidden="true">⌕</span><input id="busca" type="search" placeholder="{L('Search rules, screens, tasks…')}" aria-label="{L('Search the guide')}" autocomplete="off"></label>
    {idiomas}
    <a class="btn-topo" href="{raiz}{TOPO}../?demo#/entrar">{L('Open the prototype')}</a>
    <button class="menu-celular" type="button" aria-label="{L('Menu and search')}">☰</button>
  </div>
  <div class="busca-resultados" id="busca-resultados" role="listbox" hidden></div>
</header>
"""


def rodape(raiz):
    return f"""<footer class="rodape">
  <div><b>KORbuild</b> · {L('KORbuild product guide')} · {L('updated')} {hoje()} · <a href="{raiz}{TOPO}../">{L('Main sign-in')}</a> · <a href="{raiz}{TOPO}../?demo#/entrar">{L('Prototype')}</a></div>
  <div>{L('Screenshots come from the prototype with sample data. Names, companies and numbers are fictitious.')}</div>
</footer>
<button class="voltar-topo" type="button" aria-label="{L('Back to top')}">↑</button>
<script src="{raiz}{TOPO}assets/portal.js"></script>
</body></html>
"""


def toc_html(toc):
    if not toc:
        return '<span></span>'
    itens = []
    for t in toc:
        itens.append(f'<li><a href="#{t["id"]}">{E(html.unescape(t["name"]))}</a></li>')
        itens += [f'<li class="sub"><a href="#{c["id"]}">{E(html.unescape(c["name"]))}</a></li>' for c in t['children'][:14]]
    return f'<nav class="toc" aria-label="{L("On this page")}"><div class="toc-titulo">{L("On this page")}</div><ul>{"".join(itens)}</ul></nav>'


# ---------------------------------------------------------------- páginas

def pagina_documento(mod, p, corpo, toc, palavra_word):
    raiz = '../'
    lat = []
    for g in GRUPOS:
        itens = [x for x in mod['paginas'] if x['grupo'] == g]
        if itens:
            lat.append(f'<div class="lat-titulo">{L(g)}</div><ul>' + ''.join(
                f'<li><a href="{x["slug"]}.html"{" aria-current=page" if x is p else ""}>{E(x["titulo"])}</a></li>' for x in itens) + '</ul>')
    return cabecalho(f"{p['titulo']} · {L(mod['nome'])} · {L('KORbuild guide')}", raiz, 'modulos', f"{mod['id']}/{p['slug']}.html") + f"""
<div class="doc-layout" style="--cor:{mod['cor']}">
  <aside class="lateral" aria-label="{L(mod['nome'])} {L('documents')}">
    <a class="lat-modulo" href="index.html"><span class="ico">{mod['ico']}</span>{L(mod['nome'])}</a>
    {''.join(lat)}
  </aside>
  <main class="doc" id="conteudo">
    <nav class="trilha" aria-label="{L('You are here')}"><a href="{raiz}index.html">{L('Home')}</a> › <a href="index.html">{L(mod['nome'])}</a> › {E(L(p['grupo']))}</nav>
    <header class="doc-cabecalho">
      <h1>{E(p['titulo'])}</h1>
      <div class="doc-acoes"><a class="botao" href="{raiz}files/{palavra_word}" download>⤓ {L('Download Word')}</a><button class="botao botao-secundario" type="button" onclick="window.print()">⎙ {L('Print')}</button></div>
    </header>
    <article class="texto">
{corpo}
    </article>
  </main>
  {toc_html(toc)}
</div>
""" + rodape(raiz)


def pagina_geral(g, corpo, toc, palavra_word):
    raiz = ''
    return cabecalho(f"{g['titulo']} · {L('KORbuild guide')}", raiz, g['slug'], f"{g['slug']}.html") + f"""
<div class="doc-layout" style="--cor:var(--cor-geral)">
  <aside class="lateral" aria-label="{L('Modules')}">
    <a class="lat-modulo" href="index.html"><span class="ico">K</span>KORbuild</a>
    <div class="lat-titulo">{L('Guide')}</div><ul>{''.join(f'<li><a href="{x["slug"]}.html"{" aria-current=page" if x is g else ""}>{E(x["titulo"])}</a></li>' for x in GERAIS)}</ul>
    <div class="lat-titulo">{L('Modules')}</div><ul>{''.join(f'<li><a href="{m["id"]}/index.html">{L(m["nome"])}</a></li>' for m in MODULOS)}</ul>
  </aside>
  <main class="doc" id="conteudo">
    <nav class="trilha" aria-label="{L('You are here')}"><a href="index.html">{L('Home')}</a> › {L('Guide')}</nav>
    <header class="doc-cabecalho">
      <h1>{E(g['titulo'])}</h1>
      <div class="doc-acoes"><a class="botao" href="files/{palavra_word}" download>⤓ {L('Download Word')}</a><button class="botao botao-secundario" type="button" onclick="window.print()">⎙ {L('Print')}</button></div>
    </header>
    <article class="texto">
{corpo}
    </article>
  </main>
  {toc_html(toc)}
</div>
""" + rodape(raiz)


def pagina_modulo(mod):
    raiz = '../'
    grupos = []
    for g in GRUPOS:
        itens = [x for x in mod['paginas'] if x['grupo'] == g]
        if itens:
            grupos.append(f'<h2>{L(g)}</h2><div class="cartoes">' + ''.join(
                f'<a class="cartao-doc" href="{x["slug"]}.html"><b>{E(x["titulo"])}</b><span>{E(L(x["desc"]))}</span></a>' for x in itens) + '</div>')
    if not grupos:
        grupos.append(f'<h2>{L("Coming soon")}</h2><p>{L("The documentation of this module is being written. Meanwhile, see it working in the")} '
                      f'<a href="{raiz}{TOPO}../?demo#/entrar">{L("prototype")}</a>.</p>')
    return cabecalho(f"{L(mod['nome'])} · {L('KORbuild guide')}", raiz, 'modulos', f"{mod['id']}/index.html") + f"""
<div class="mod-cabeca" style="--cor:{mod['cor']}">
  <nav class="trilha" aria-label="{L('You are here')}"><a href="{raiz}index.html">{L('Home')}</a> › {L(mod['nome'])}</nav>
  <div class="mod-caixa"><span class="ico">{mod['ico']}</span><div><h1>{L(mod['nome'])}</h1><p>{E(L(mod['resumo']))}</p></div></div>
</div>
<div class="mod-grupos" style="--cor:{mod['cor']}">{''.join(grupos)}</div>
""" + rodape(raiz)


def pagina_inicial():
    cartoes = ''.join(
        f'<a class="cartao-mod" href="{m["id"]}/index.html" style="--cor:{m["cor"]}"><span class="ico">{m["ico"]}</span><b>{L(m["nome"])}</b>'
        f'<span>{E(L(m["resumo"]))}</span>' + ('' if m['paginas'] else f'<span class="em-breve">{L("Documentation coming soon")}</span>') + '</a>'
        for m in MODULOS)
    prot = TOPO + '../?demo#/entrar'
    lista = lambda pares: '<ul>' + ''.join(f'<li><a href="{h}">{t}</a></li>' for h, t in pares) + '</ul>'
    mods = [L(m['nome']) for m in MODULOS]
    return cabecalho(L('KORbuild product guide'), '', 'inicio', 'index.html') + f"""
<section class="heroi"><div class="heroi-dentro">
  <p class="sobrescrito">{L('KORbuild product guide')}</p>
  <h1>{L('Everything KORbuild does, in one place.')}</h1>
  <p class="lead">{L('KORbuild is a platform for construction companies in the United States: general contractors and the trade contractors they hire. This guide explains each module — what it is for, the rules it enforces, who can do what — and walks through every task with screenshots, so anyone can present the product with confidence.')}</p>
  <div class="botoes">
    <a class="botao" href="presenting.html">{L("Presenter's guide →")}</a>
    <a class="botao botao-secundario" href="organization.html">{L('How KORbuild is organized')}</a>
    <a class="botao botao-secundario" href="#modules">{L('Explore the modules')}</a>
    <a class="botao botao-secundario" href="{prot}">{L('Open the prototype ↗')}</a>
  </div>
</div></section>
<section class="faixa" id="modules">
  <h2>{L('Modules')}</h2>
  <p>{L('One account, one set of jobs and people, shared by every module. Each company turns on the modules it needs.')}</p>
  <div class="cartoes">{cartoes}</div>
</section>
<section class="faixa">
  <h2>{L('Where to start')}</h2>
  <p>{L('Pick the path that matches why you are here.')}</p>
  <div class="publicos">
    <div class="publico"><h3>{L('I will present KORbuild to a customer')}</h3><p>{L("Read the presenter's guide: the story, the order of the screens and what to say.")}</p>
      {lista([('organization.html', L('How KORbuild is organized')), ('presenting.html', L("Presenter's guide")), (prot, L('Open the prototype'))])}</div>
    <div class="publico"><h3>{L('I will use KORbuild every day')}</h3><p>{L('Go to your module and open the user manual: one task at a time, with screenshots.')}</p>
      {lista([(f"{m['id']}/user-manual.html", f"{L(m['nome'])} — {L('user manual')}") for m in MODULOS])}</div>
    <div class="publico"><h3>{L('I need the rules and permissions')}</h3><p>{L('Each module lists its business rules with IDs and who can do what.')}</p>
      {lista([(f"{m['id']}/{pg}.html", f"{L(m['nome'])} — {L(rot)}") for m in MODULOS for pg, rot in (('business-rules', 'business rules'), ('roles-and-permissions', 'roles and permissions'))])}</div>
  </div>
</section>
<div style="height:40px"></div>
""" + rodape('')


# ---------------------------------------------------------------- busca

class Texto(HTMLParser):
    def __init__(self):
        super().__init__(); self.partes = []
    def handle_data(self, d): self.partes.append(d)


def texto_puro(trecho):
    t = Texto(); t.feed(trecho)
    return re.sub(r'\s+', ' ', ''.join(t.partes)).strip()


def entradas_de_busca(url, pagina, corpo):
    saida = []
    pedacos = re.split(r'(<h[23] id="[^"]+">.*?</h[23]>)', corpo)
    titulo, ancora = pagina, ''
    for pedaco in pedacos:
        m = re.match(r'<h[23] id="([^"]+)">(.*?)</h[23]>', pedaco)
        if m:
            ancora, titulo = m.group(1), texto_puro(m.group(2))
            continue
        txt = texto_puro(pedaco)
        if txt:
            saida.append({'url': url + ('#' + ancora if ancora else ''), 'pagina': pagina, 'titulo': titulo, 'texto': txt[:600]})
    return saida


# ---------------------------------------------------------------- Word

def gerar_word(titulo, subtitulo, corpo_html, destino, raiz_img):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # US Letter
    for lado in ('left_margin', 'right_margin'):
        setattr(sec, lado, Cm(2.2))
    estilo = doc.styles['Normal']; estilo.font.name = 'Calibri'; estilo.font.size = Pt(10.5)
    for n, tam in ((1, 18), (2, 14), (3, 12)):
        h = doc.styles[f'Heading {n}']; h.font.name = 'Calibri'; h.font.size = Pt(tam); h.font.color.rgb = RGBColor(0x14, 0x1B, 0x26)
    p = doc.add_paragraph(); r = p.add_run('KORbuild'); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(0xC2, 0x41, 0x0C)
    doc.add_heading(titulo, 0)
    if subtitulo:
        doc.add_paragraph(subtitulo).runs[0].italic = True

    def sombrear(par, cor):
        ppr = par._p.get_or_add_pPr(); shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), cor); ppr.append(shd)

    numbering = doc.part.numbering_part.numbering_definitions._numbering
    estilo_num = doc.styles['List Number']._element.pPr.numPr.numId.val
    abstrato = [n for n in numbering.findall(qn('w:num')) if n.get(qn('w:numId')) == str(estilo_num)][0].find(qn('w:abstractNumId')).get(qn('w:val'))

    def nova_numeracao():
        ids = [int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]
        novo = max(ids) + 1
        num = OxmlElement('w:num'); num.set(qn('w:numId'), str(novo))
        ab = OxmlElement('w:abstractNumId'); ab.set(qn('w:val'), abstrato); num.append(ab)
        ov = OxmlElement('w:lvlOverride'); ov.set(qn('w:ilvl'), '0')
        st = OxmlElement('w:startOverride'); st.set(qn('w:val'), '1'); ov.append(st); num.append(ov)
        numbering.append(num)
        return novo

    def numerar(par, num_id, nivel):
        ppr = par._p.get_or_add_pPr()
        npr = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl'); il.set(qn('w:val'), str(nivel)); npr.append(il)
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(num_id)); npr.append(ni)
        ppr.append(npr)

    CORES = {'who': 'ECFDF3', 'warning': 'FFFAEB', 'simulated': 'FFFAEB', 'tip': 'FFF2EA', 'confirm': 'FEF3F2', 'say': 'EEF1F5', 'note': 'EEF4FF'}

    class Montar(HTMLParser):
        def __init__(self):
            super().__init__()
            self.par = None; self.fmt = []; self.listas = []; self.tabela = None; self.linha = None; self.celula = None
            self.cabecalho = False; self.caixa = None; self.ignorar = 0; self.figcaption = False; self.numeros = []
        def novo(self, estilo=None):
            alvo = self.celula if self.celula is not None else doc
            self.par = alvo.add_paragraph(style=estilo) if estilo else alvo.add_paragraph()
            if self.caixa:
                sombrear(self.par, CORES.get(self.caixa, 'EEF4FF'))
            return self.par
        def handle_starttag(self, tag, attrs):
            a = dict(attrs)
            if tag in ('h1', 'h2', 'h3', 'h4'):
                self.par = doc.add_heading('', level=min(3, int(tag[1]) - 1 if tag != 'h1' else 1))
            elif tag == 'p':
                classe = a.get('class', '')
                if not (self.par is not None and not self.par.text and self.listas):
                    self.novo()
                if 'admonition-title' in classe:
                    self.fmt.append('b')
            elif tag in ('ul', 'ol'):
                self.listas.append(tag)
                self.numeros.append(nova_numeracao() if tag == 'ol' else None)
            elif tag == 'li':
                nivel = len(self.listas)
                base = 'List Number' if self.listas and self.listas[-1] == 'ol' else 'List Bullet'
                self.novo(base + (f' {nivel}' if nivel > 1 else ''))
                if self.numeros and self.numeros[-1] is not None:
                    numerar(self.par, self.numeros[-1], nivel - 1)
            elif tag in ('strong', 'b'):
                self.fmt.append('b')
            elif tag in ('em', 'i'):
                self.fmt.append('i')
            elif tag == 'code':
                self.fmt.append('c')
            elif tag == 'div' and 'admonition' in a.get('class', ''):
                self.caixa = (a.get('class', '').split() + ['note'])[1]
            elif tag == 'table':
                self.tabela = doc.add_table(rows=0, cols=0); self.tabela.style = 'Table Grid'; self.colunas = 0
            elif tag == 'tr':
                self.linha = []
            elif tag in ('td', 'th'):
                self.cabecalho = tag == 'th'; self.linha.append([]); self.em_celula = True
            elif tag == 'img':
                arq = raiz_img / a.get('src', '').split('img/', 1)[-1]
                if arq.exists():
                    largura = Cm(7) if '-mobile' in arq.name else Cm(16)
                    doc.add_picture(str(arq), width=largura)
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif tag == 'figcaption':
                self.figcaption = True; self.novo(); self.par.alignment = WD_ALIGN_PARAGRAPH.CENTER; self.fmt.append('i')
            elif tag == 'br' and self.par is not None:
                self.par.add_run().add_break()
        def handle_endtag(self, tag):
            if tag in ('strong', 'b', 'em', 'i', 'code') and self.fmt:
                self.fmt.pop()
            elif tag == 'p' and self.fmt and self.fmt[-1] == 'b' and self.par is not None and self.caixa:
                self.fmt.pop()
            elif tag in ('ul', 'ol') and self.listas:
                self.listas.pop()
                if self.numeros: self.numeros.pop()
            elif tag == 'div' and self.caixa:
                self.caixa = None
            elif tag in ('td', 'th'):
                self.em_celula = False
            elif tag == 'tr' and self.linha is not None:
                linhas = self._linhas = getattr(self, '_linhas', []); linhas.append((self.linha, self.cabecalho))
            elif tag == 'table':
                linhas = getattr(self, '_linhas', [])
                if linhas:
                    ncol = max(len(l) for l, _ in linhas)
                    t = doc.add_table(rows=len(linhas), cols=ncol); t.style = 'Table Grid'
                    for i, (cels, cab) in enumerate(linhas):
                        for j, partes in enumerate(cels):
                            cel = t.cell(i, j); cel.text = ''
                            par = cel.paragraphs[0]
                            for txt, f in partes:
                                run = par.add_run(txt); run.bold = 'b' in f or (i == 0); run.italic = 'i' in f; run.font.size = Pt(9)
                            if i == 0:
                                sombrear(par, 'F2F4F7')
                    # remove a tabela vazia criada no início
                    self.tabela._tbl.getparent().remove(self.tabela._tbl)
                self.tabela = None; self._linhas = []
            elif tag == 'figcaption':
                self.figcaption = False
                if self.fmt: self.fmt.pop()
        def handle_data(self, dado):
            if not dado.strip():
                if '\n' in dado or self.par is None or not self.par.text:
                    return
                dado = ' '
            if getattr(self, 'em_celula', False):
                self.linha[-1].append((dado, set(self.fmt)))
                return
            if self.tabela is not None:
                return
            if self.par is None:
                self.novo()
            run = self.par.add_run(dado)
            run.bold = 'b' in self.fmt; run.italic = 'i' in self.fmt
            if 'c' in self.fmt:
                run.font.name = 'Consolas'

    Montar().feed(corpo_html)
    doc.save(destino)


# ---------------------------------------------------------------- conferência de links

def conferir_links():
    erros = []
    for pagina in PORTAL.rglob('*.html'):
        if 'fonte' in pagina.parts:
            continue
        txt = pagina.read_text(encoding='utf-8')
        ids = set(re.findall(r'id="([^"]+)"', txt))
        for alvo in re.findall(r'(?:href|src)="([^"]+)"', txt):
            if re.match(r'(https?:|mailto:|data:)', alvo) or '?demo' in alvo:
                continue
            caminho, _, ancora = alvo.partition('#')
            if not caminho:
                if ancora and ancora not in ids:
                    erros.append(f'{pagina.relative_to(RAIZ)}: âncora #{ancora} não existe')
                continue
            arq = (pagina.parent / caminho).resolve()
            if arq.is_dir():
                arq = arq / 'index.html'
            if not arq.exists():
                erros.append(f'{pagina.relative_to(RAIZ)}: {alvo} não existe')
            elif ancora and arq.suffix == '.html' and f'id="{ancora}"' not in arq.read_text(encoding='utf-8'):
                erros.append(f'{pagina.relative_to(RAIZ)}: {alvo} (âncora não existe)')
    return erros


# ---------------------------------------------------------------- principal

def gerar_idioma(idioma):
    definir_idioma(idioma)
    origem = FONTE / idioma
    (SAIDA / 'files').mkdir(parents=True, exist_ok=True)
    busca, faltando = [], []
    sub = L('KORbuild product guide')
    # os títulos de todas as páginas antes de desenhar (o menu lateral mostra os das outras páginas)
    for g in GERAIS:
        arq = origem / f'{g["slug"]}.md'
        if arq.exists():
            g['titulo'] = titulo_e_corpo(arq)[0]
    for mod in MODULOS:
        for p in mod['paginas']:
            arq = origem / mod['id'] / f'{p["slug"]}.md'
            if arq.exists():
                p['titulo'] = titulo_e_corpo(arq)[0]
    for g in GERAIS:
        arq = origem / f'{g["slug"]}.md'
        if not arq.exists():
            faltando.append(str(arq)); continue
        titulo, md = titulo_e_corpo(arq)
        corpo, toc = converter(md, '')
        word = f'KORbuild - {titulo}.docx'.replace("'", '')
        (SAIDA / f'{g["slug"]}.html').write_text(pagina_geral(g, corpo, toc, word), encoding='utf-8')
        gerar_word(titulo, sub, corpo, SAIDA / 'files' / word, SAIDA / 'img')
        busca += entradas_de_busca(f'{g["slug"]}.html', titulo, corpo)
    for mod in MODULOS:
        pasta = SAIDA / mod['id']; pasta.mkdir(exist_ok=True)
        for p in mod['paginas']:
            arq = origem / mod['id'] / f'{p["slug"]}.md'
            if not arq.exists():
                faltando.append(str(arq)); continue
            titulo, md = titulo_e_corpo(arq)
            corpo, toc = converter(md, '../')
            word = f'KORbuild {L(mod["nome"])} - {titulo}.docx'
            (pasta / f'{p["slug"]}.html').write_text(pagina_documento(mod, p, corpo, toc, word), encoding='utf-8')
            gerar_word(f'{L(mod["nome"])} — {titulo}', sub, corpo, SAIDA / 'files' / word, SAIDA / 'img')
            busca += entradas_de_busca(f'{mod["id"]}/{p["slug"]}.html', f'{L(mod["nome"])} · {titulo}', corpo)
        (pasta / 'index.html').write_text(pagina_modulo(mod), encoding='utf-8')
    (SAIDA / 'index.html').write_text(pagina_inicial(), encoding='utf-8')
    (SAIDA / 'search.json').write_text(json.dumps(busca, ensure_ascii=False), encoding='utf-8')
    return busca, faltando


def main():
    so = [a.split('=', 1)[1] for a in sys.argv[1:] if a.startswith('--idioma=')]
    total, faltando = {}, []
    for idioma in (so or IDIOMAS):
        busca, falta = gerar_idioma(idioma)
        total[idioma] = len(busca); faltando += falta
    erros = conferir_links()
    for f in faltando:
        print('falta a fonte:', f)
    for e in erros:
        print('LINK QUEBRADO', e)
    print('portal gerado: ' + ', '.join(f'{k} {v} trechos na busca' for k, v in total.items()) + f', {len(erros)} links quebrados')
    sys.exit(1 if erros or faltando else 0)


if __name__ == '__main__':
    main()
